#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script para analizar bordes y estilos detallados del archivo BSM original
"""

import openpyxl
from openpyxl.utils import get_column_letter

# Abrir el archivo Excel original de BSM
archivo = '/Users/albertosanmartin/Documents/VSS-COTIZACIONES/brokers28-12-2025/BSM CATERING/RFQ_EQ_SCF_BAA_0031.xlsm'
wb = openpyxl.load_workbook(archivo, data_only=True)
ws = wb.active

print(f"Archivo: {archivo}")
print("="*100)

# Analizar las primeras 18 filas para ver bordes y líneas
print("\n### BORDES Y LÍNEAS EN METADATA ###\n")

for row_idx in range(1, 19):
    print(f"\n--- FILA {row_idx} ---")
    
    # Revisar todas las columnas hasta la V (22)
    cells_with_borders = []
    for col_idx in range(1, 23):
        cell = ws.cell(row=row_idx, column=col_idx)
        col_letter = get_column_letter(col_idx)
        
        # Revisar si tiene bordes
        if cell.border:
            border_info = []
            if cell.border.left and cell.border.left.style:
                border_info.append(f"Left: {cell.border.left.style} ({cell.border.left.color})")
            if cell.border.right and cell.border.right.style:
                border_info.append(f"Right: {cell.border.right.style} ({cell.border.right.color})")
            if cell.border.top and cell.border.top.style:
                border_info.append(f"Top: {cell.border.top.style} ({cell.border.top.color})")
            if cell.border.bottom and cell.border.bottom.style:
                border_info.append(f"Bottom: {cell.border.bottom.style} ({cell.border.bottom.color})")
            
            if border_info:
                cells_with_borders.append(f"{col_letter}{row_idx}: {', '.join(border_info)}")
    
    if cells_with_borders:
        for border_str in cells_with_borders:
            print(f"  {border_str}")

# Buscar imágenes en el worksheet
print("\n\n### IMÁGENES EN EL WORKSHEET ###\n")
if hasattr(ws, '_images') and ws._images:
    for idx, img in enumerate(ws._images):
        print(f"Imagen {idx + 1}:")
        try:
            if hasattr(img.anchor, '_from'):
                print(f"  Desde: Col {img.anchor._from.col}, Row {img.anchor._from.row}")
            if hasattr(img.anchor, 'to'):
                print(f"  Hasta: Col {img.anchor.to.col}, Row {img.anchor.to.row}")
            print(f"  Width: {img.width}, Height: {img.height}")
        except Exception as e:
            print(f"  Error al analizar anchor: {e}")
        if hasattr(img, 'path'):
            print(f"  Path: {img.path}")
else:
    print("No se encontraron imágenes en el worksheet")

# Analizar filas con más detalle de fill
print("\n\n### FILLS Y BACKGROUNDS DETALLADOS ###\n")
for row_idx in [3, 12, 18, 19]:
    print(f"\n--- FILA {row_idx} ---")
    for col_idx in range(1, 23):
        cell = ws.cell(row=row_idx, column=col_idx)
        col_letter = get_column_letter(col_idx)
        
        if cell.fill and cell.fill.fill_type:
            fill_type = cell.fill.fill_type
            if fill_type == 'solid':
                fg_color = cell.fill.fgColor
                color_str = str(fg_color.rgb) if hasattr(fg_color, 'rgb') else str(fg_color.index)
                if color_str != '00000000':  # No es transparente
                    print(f"  {col_letter}{row_idx}: Fill={color_str}, Type={fill_type}")

wb.close()
print("\n✅ Análisis completado")
