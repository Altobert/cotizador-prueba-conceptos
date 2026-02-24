#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script para analizar la estructura del archivo Excel BSM original
"""

import openpyxl
from openpyxl.utils import get_column_letter

# Abrir el archivo Excel original de BSM
archivo = '/Users/albertosanmartin/Documents/VSS-COTIZACIONES/brokers28-12-2025/BSM CATERING/RFQ_EQ_SCF_BAA_0031.xlsm'
wb = openpyxl.load_workbook(archivo, data_only=True)
ws = wb.active

print(f"Archivo: {archivo}")
print(f"Hoja activa: {ws.title}")
print(f"Dimensiones: {ws.dimensions}")
print("="*100)

# Analizar las primeras 25 filas para ver la estructura de metadata
print("\n### ESTRUCTURA DE METADATA Y HEADER ###\n")

for row_idx in range(1, 26):
    print(f"\n--- FILA {row_idx} ---")
    
    # Revisar todas las columnas hasta la V (22)
    for col_idx in range(1, 23):
        cell = ws.cell(row=row_idx, column=col_idx)
        col_letter = get_column_letter(col_idx)
        
        if cell.value is not None:
            # Información de la celda
            valor = str(cell.value)[:60]  # Primeros 60 caracteres
            
            # Información de estilo
            bg_color = "None"
            if cell.fill and cell.fill.fgColor:
                bg_color = str(cell.fill.fgColor.rgb) if hasattr(cell.fill.fgColor, 'rgb') else str(cell.fill.fgColor.index)
            
            font_info = ""
            if cell.font:
                font_info = f"Font: {cell.font.name or 'default'}, Size: {cell.font.size or 'default'}"
                if cell.font.bold:
                    font_info += ", BOLD"
                if cell.font.color:
                    font_color = str(cell.font.color.rgb) if hasattr(cell.font.color, 'rgb') else str(cell.font.color.index)
                    font_info += f", Color: {font_color}"
            
            align_info = ""
            if cell.alignment:
                if cell.alignment.horizontal:
                    align_info += f"H:{cell.alignment.horizontal} "
                if cell.alignment.vertical:
                    align_info += f"V:{cell.alignment.vertical}"
            
            print(f"  {col_letter}{row_idx}: '{valor}'")
            print(f"    BG: {bg_color} | {font_info}")
            if align_info:
                print(f"    Align: {align_info}")
            
            # Información de merged cells
            if cell.coordinate in ws.merged_cells:
                print(f"    ⚠️  Celda mezclada")

# Información de celdas mezcladas
print("\n\n### CELDAS MEZCLADAS ###\n")
for merged_range in ws.merged_cells.ranges:
    print(f"  {merged_range}")

# Analizar anchos de columna
print("\n\n### ANCHOS DE COLUMNA ###\n")
for col_idx in range(1, 23):
    col_letter = get_column_letter(col_idx)
    if col_letter in ws.column_dimensions:
        width = ws.column_dimensions[col_letter].width
        print(f"  Columna {col_letter}: {width}")

# Analizar alturas de fila para las primeras 25 filas
print("\n\n### ALTURAS DE FILA (1-25) ###\n")
for row_idx in range(1, 26):
    if row_idx in ws.row_dimensions:
        height = ws.row_dimensions[row_idx].height
        if height:
            print(f"  Fila {row_idx}: {height}")

wb.close()
print("\n✅ Análisis completado")
